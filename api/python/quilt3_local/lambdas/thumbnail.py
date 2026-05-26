from __future__ import annotations

import contextlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import urllib.parse
from io import BytesIO

import requests
from PIL import Image

from .._upstream import load_module
from .shared.decorator import QUILT_INFO_HEADER, api, validate
from .shared.utils import get_default_origins, make_json_response

_upstream = load_module("lambdas.thumbnail")

DEFAULT_PDF_RENDER_DPI = 300
MAX_PDF_RENDER_DPI = 300

SIZE_PARAMETER_MAP = getattr(
	_upstream,
	"SIZE_PARAMETER_MAP",
	{
		"w32h32": (32, 32),
		"w64h64": (64, 64),
		"w128h128": (128, 128),
		"w256h256": (256, 256),
		"w480h320": (480, 320),
		"w640h480": (640, 480),
		"w960h640": (960, 640),
		"w1024h768": (1024, 768),
		"w2048h1536": (2048, 1536),
	},
)

SCHEMA = {
	"type": "object",
	"properties": {
		"url": {"type": "string"},
		"size": {"enum": list(SIZE_PARAMETER_MAP)},
		"input": {"enum": ["pdf", "pptx"]},
		"page": {"type": "string", "pattern": r"^\d+$"},
		"countPages": {"enum": ["true", "false"]},
	},
	"required": ["url", "size"],
	"additionalProperties": False,
}


class PDFThumbError(Exception):
	pass


def _run_command(*argv: str) -> subprocess.CompletedProcess[str]:
	try:
		return subprocess.run(argv, check=True, capture_output=True, text=True)
	except FileNotFoundError as exc:
		raise PDFThumbError(f"Missing required command: {argv[0]}") from exc
	except subprocess.CalledProcessError as exc:
		detail = (exc.stderr or exc.stdout or str(exc)).strip()
		raise PDFThumbError(detail) from exc


def _render_pdf_page_with_pdfium(*, path: str, page: int, dpi: int) -> Image.Image:
	try:
		import pypdfium2
	except ImportError as exc:
		raise PDFThumbError("Missing required dependency: pypdfium2") from exc

	document = pypdfium2.PdfDocument(path)
	page_index = page - 1
	if page_index < 0 or page_index >= len(document):
		raise PDFThumbError(f"Page {page} is out of range for {path}")
	bitmap = document[page_index].render(scale=dpi / 72)
	return bitmap.to_pil()


def _count_pdf_pages_with_pdfium(path: str) -> int:
	try:
		import pypdfium2
	except ImportError as exc:
		raise PDFThumbError("Missing required dependency: pypdfium2") from exc
	return len(pypdfium2.PdfDocument(path))


def get_pdf_render_dpi() -> int:
	raw = os.environ.get("PDF_PREVIEW_DPI")
	if raw is None:
		return DEFAULT_PDF_RENDER_DPI
	try:
		dpi = int(raw)
	except ValueError as exc:
		raise PDFThumbError(f"Invalid PDF_PREVIEW_DPI: {raw!r}") from exc
	return max(72, min(dpi, MAX_PDF_RENDER_DPI))


def resize_pdf_page(img: Image.Image, *, size: int) -> Image.Image:
	if img.width <= size:
		return img

	height = max(1, round(img.height * size / img.width))
	return img.resize((size, height), Image.Resampling.LANCZOS)


def render_pdf_page(*, path: str, page: int, dpi: int) -> Image.Image:
	if shutil.which("pdftoppm") is None:
		return _render_pdf_page_with_pdfium(path=path, page=page, dpi=dpi)

	with tempfile.TemporaryDirectory() as out_dir:
		out_base = os.path.join(out_dir, "page")
		_run_command(
			"pdftoppm",
			"-f",
			str(page),
			"-l",
			str(page),
			"-r",
			str(dpi),
			"-singlefile",
			"-jpeg",
			path,
			out_base,
		)
		rendered = out_base + ".jpg"
		if not os.path.exists(rendered):
			raise PDFThumbError("pdftoppm did not produce an output image")
		with Image.open(rendered) as img:
			return img.copy()


def count_pdf_pages(path: str) -> int:
	if shutil.which("pdfinfo") is None:
		return _count_pdf_pages_with_pdfium(path)

	result = _run_command("pdfinfo", path)
	match = re.search(r"^Pages:\s+(\d+)\s*$", result.stdout, re.MULTILINE)
	if match is None:
		raise PDFThumbError("Unable to determine PDF page count")
	return int(match.group(1))


@contextlib.contextmanager
def pptx_to_pdf(*, path: str, page: int):
	with tempfile.TemporaryDirectory() as out_dir:
		with tempfile.TemporaryDirectory() as tmp_dir:
			try:
				subprocess.run(
					(
						"libreoffice",
						"--convert-to",
						'pdf:impress_pdf_Export:{"PageRange":{"type":"string","value":"%s-%s"}}'
						% (page, page),
						"--outdir",
						out_dir,
						path,
					),
					check=True,
					env={
						**os.environ,
						"HOME": tmp_dir,
					},
				)
			except FileNotFoundError as exc:
				raise PDFThumbError("Missing required command: libreoffice") from exc
		yield os.path.join(out_dir, os.path.splitext(os.path.basename(path))[0] + ".pdf")


def pdf_thumb(*, path: str, page: int, size: int):
	render_dpi = get_pdf_render_dpi()
	page_image = render_pdf_page(path=path, page=page, dpi=render_dpi)
	return resize_pdf_page(page_image, size=size), render_dpi


def handle_pdf(*, path: str, page: int, size: int, count_pages: bool):
	fmt = "JPEG"
	thumb, render_dpi = pdf_thumb(path=path, page=page, size=size)
	info = {
		"thumbnail_format": fmt,
		"thumbnail_size": thumb.size,
		"pdf_render_dpi": render_dpi,
		"pdf_resize_filter": "LANCZOS",
	}
	if count_pages:
		info["page_count"] = count_pdf_pages(path)

	thumbnail_bytes = BytesIO()
	thumb.save(thumbnail_bytes, fmt)
	return info, thumbnail_bytes.getvalue()


def handle_pptx(*, path: str, page: int, size: int, count_pages: bool):
	try:
		import pptx
	except ImportError as exc:
		raise PDFThumbError("Missing required dependency: python-pptx") from exc

	with pptx_to_pdf(path=path, page=page) as pdf_path:
		info, data = handle_pdf(path=pdf_path, page=1, size=size, count_pages=False)
	if count_pages:
		info["page_count"] = len(pptx.Presentation(path).slides)
	return info, data


@api(cors_origins=get_default_origins())
@validate(SCHEMA)
def lambda_handler(request):
	input_ = request.args.get("input")
	if input_ not in {"pdf", "pptx"}:
		return _upstream.lambda_handler(request.event, request.context)

	url = request.args["url"]
	size = SIZE_PARAMETER_MAP[request.args["size"]]
	page = int(request.args.get("page", "1"))
	count_pages = request.args.get("countPages") == "true"

	resp = requests.get(url)
	if not resp.ok:
		return make_json_response(resp.status_code, {"error": resp.reason, "text": resp.text})

	filename_suffix = urllib.parse.unquote(urllib.parse.urlparse(url).path.split("/")[-1])
	with tempfile.NamedTemporaryFile(suffix=filename_suffix) as src_file:
		src_file.write(resp.content)
		src_file.flush()

		try:
			if input_ == "pdf":
				info, data = handle_pdf(path=src_file.name, page=page, size=size[0], count_pages=count_pages)
			else:
				info, data = handle_pptx(path=src_file.name, page=page, size=size[0], count_pages=count_pages)
		except PDFThumbError as exc:
			return make_json_response(500, {"error": str(exc)})

	headers = {
		"Content-Type": Image.MIME[info["thumbnail_format"]],
		QUILT_INFO_HEADER: json.dumps(info),
	}
	return 200, data, headers
